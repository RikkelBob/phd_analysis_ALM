import numpy as np
import os
import yaml
import config
import matplotlib.pyplot as plt
import matplotlib as mpl
from neural_analysis import plot_utils
from statsmodels.stats.multitest import multipletests

data_dir = ""  # dir containing dpca results for a single run

# ============================================================================
# PAPER FORMAT SETTINGS
# ============================================================================
SUBPLOT_W = 0.9; SUBPLOT_H = 0.75
H_PAD = 0.9; V_PAD = 0.7
LEFT_MARGIN = 0.7; RIGHT_MARGIN = 0.3; TOP_MARGIN = 0.6; BOTTOM_MARGIN = 0.5
MM_TO_IN = 1 / 25.4
FONT_DEFAULT = 6; FONT_SUPTITLE = 10; FONT_LEGEND = 6; FONT_TITLE = 8
X_LIM_DEFAULT = [-0.5, 1.5]

mpl.rcParams.update({
    'font.size': FONT_DEFAULT, 'axes.titlesize': FONT_TITLE, 'axes.labelsize': FONT_DEFAULT,
    'xtick.labelsize': FONT_DEFAULT, 'ytick.labelsize': FONT_DEFAULT, 'legend.fontsize': FONT_LEGEND,
    'figure.titlesize': FONT_SUPTITLE, 'axes.linewidth': 0.8,
    'xtick.major.width': 0.6, 'ytick.major.width': 0.6, 'xtick.major.size': 3, 'ytick.major.size': 3,
    'lines.linewidth': 1.5, 'font.family': 'sans-serif',
    'font.sans-serif': ['Arial', 'Helvetica', 'DejaVu Sans'],
    'pdf.fonttype': 42, 'ps.fonttype': 42, 'svg.fonttype': 'none',
})

# ============================================================================
# SUBPLOT HELPERS
# ============================================================================
def make_figure_with_variable_panels(panel_sizes_in):
    n = len(panel_sizes_in); max_w = max(w for w, h in panel_sizes_in)
    fig_w = LEFT_MARGIN + max_w + RIGHT_MARGIN
    fig_h = TOP_MARGIN + sum(h for w, h in panel_sizes_in) + ((n-1)*V_PAD if n > 1 else 0) + BOTTOM_MARGIN
    fig = plt.figure(figsize=(fig_w, fig_h)); axes = np.empty(n, dtype=object)
    cursor_y = fig_h - TOP_MARGIN
    for i, (pw, ph) in enumerate(panel_sizes_in):
        cursor_y -= ph
        axes[i] = fig.add_axes([LEFT_MARGIN/fig_w, cursor_y/fig_h, pw/fig_w, ph/fig_h])
        cursor_y -= V_PAD
    return fig, axes

# ============================================================================
# DATA LOADING
# ============================================================================
def load_euclidean(data_dir, marg_key):
    data = np.load(os.path.join(data_dir, f'euclidean_{marg_key}.npz'), allow_pickle=True)
    time_centers = data['time_centers']; pairs = []; i = 0
    while f'pair{i}_label' in data:
        pair = {'label': str(data[f'pair{i}_label']), 'distance': data[f'pair{i}_distance'],
                'ci_low': data[f'pair{i}_ci_low'], 'ci_high': data[f'pair{i}_ci_high'],
                'boot_distribution': data[f'pair{i}_boot_distribution'],
                'n_trials_A': int(data[f'pair{i}_n_trials_A']), 'n_trials_B': int(data[f'pair{i}_n_trials_B'])}
        if f'pair{i}_p_values' in data: pair['p_values'] = data[f'pair{i}_p_values']
        if f'pair{i}_null_distribution' in data: pair['null_distribution'] = data[f'pair{i}_null_distribution']
        # --- window-averaged test fields (if saved by analysis script) ---
        if f'pair{i}_observed_window_avg' in data:
            pair['observed_window_avg'] = float(data[f'pair{i}_observed_window_avg'])
            pair['p_value_window_avg'] = float(data[f'pair{i}_p_value_window_avg'])
            pair['window'] = tuple(data[f'pair{i}_window'])
            if f'pair{i}_null_window_avg' in data:
                pair['null_window_avg'] = data[f'pair{i}_null_window_avg']
        if f'pair{i}_cluster_mask' in data:
            pair['cluster_mask'] = data[f'pair{i}_cluster_mask'].astype(bool)
            pair['cluster_p_map'] = data[f'pair{i}_cluster_p_map']
            pair['cluster_p_value'] = data[f'pair{i}_cluster_p_value']
            pair['cluster_t_start'] = data[f'pair{i}_cluster_t_start']
            pair['cluster_t_end'] = data[f'pair{i}_cluster_t_end']
            pair['cluster_mass'] = data[f'pair{i}_cluster_mass']
        pairs.append(pair); i += 1
    return time_centers, pairs

def load_euclidean_diff(data_dir, marg_key):
    data = np.load(os.path.join(data_dir, f'euclidean_diff_{marg_key}.npz'), allow_pickle=True)
    time_centers = data['time_centers']; diffs = []; i = 0
    while f'diff{i}_label' in data:
        diffs.append({'label': str(data[f'diff{i}_label']), 'diff_distance': data[f'diff{i}_diff_distance'],
            'diff_ci_low': data[f'diff{i}_diff_ci_low'], 'diff_ci_high': data[f'diff{i}_diff_ci_high'],
            'boot_diff': data[f'diff{i}_boot_diff'], 'frac_positive': data[f'diff{i}_frac_positive']}); i += 1
    return time_centers, diffs

def compute_diff_from_pairs(pair_ctrl, pair_opto, ci=95):
    alpha = (100-ci)/2; dd = pair_ctrl['distance'] - pair_opto['distance']
    bd = pair_ctrl['boot_distribution'] - pair_opto['boot_distribution']
    return {'diff_distance': dd, 'diff_ci_low': np.nanpercentile(bd, alpha, axis=0),
            'diff_ci_high': np.nanpercentile(bd, 100-alpha, axis=0), 'boot_diff': bd}

def build_label_to_index(pairs):
    # When labels are duplicated (e.g. 'Saliency scaling (visual hits)' appears twice),
    # the existing logic keeps the LAST occurrence. Preserve that behavior.
    return {p['label']: i for i, p in enumerate(pairs)}

def find_sig_clusters(rejected, time_centers):
    clusters = []; in_c = False
    for i, s in enumerate(rejected):
        if s and not in_c: start = i; in_c = True
        elif not s and in_c: clusters.append((start, i-1)); in_c = False
    if in_c: clusters.append((start, len(rejected)-1))
    return [(time_centers[s], time_centers[e]) for s, e in clusters]

def summarize_significance(time_centers, p_values, distances, alpha=0.05):
    if np.all(np.isnan(distances)) or np.all(np.isnan(p_values)):
        return {'significant': False, 'n_sig_timepoints': 0, 'skipped': True}
    rejected, p_corr, _, _ = multipletests(p_values, alpha=alpha, method='fdr_bh')
    pre_mask = time_centers < 0; post_mask = time_centers > 0; valid = ~np.isnan(distances)
    peak_distance = round(float(np.nanmax(distances)), 6) if np.any(valid) else None
    peak_time_s = round(float(time_centers[np.nanargmax(distances)]), 4) if np.any(valid) else None
    summary = {'significant': bool(np.any(rejected)), 'n_sig_timepoints': int(np.sum(rejected)),
        'n_total_timepoints': int(len(rejected)), 'alpha': alpha, 'correction_method': 'fdr_bh',
        'mean_distance': round(float(np.nanmean(distances)), 6) if np.any(valid) else None,
        'mean_distance_pre': round(float(np.nanmean(distances[pre_mask])), 6) if np.any(pre_mask & valid) else None,
        'mean_distance_post': round(float(np.nanmean(distances[post_mask])), 6) if np.any(post_mask & valid) else None,
        'peak_distance': peak_distance, 'peak_time_s': peak_time_s,
        'min_p_uncorrected': float(np.nanmin(p_values)), 'min_p_corrected': float(np.nanmin(p_corr)),
        'n_uncorrected_sig': int(np.sum(p_values < alpha)), 'median_p_uncorrected': float(np.nanmedian(p_values))}
    if np.any(rejected):
        clusters = find_sig_clusters(rejected, time_centers)
        sig_d = distances.copy(); sig_d[~rejected] = -np.inf; pi = np.argmax(sig_d)
        summary['fdr_clusters'] = [{'start_s': round(float(s),4), 'end_s': round(float(e),4)} for s, e in clusters]
        summary['peak_sig_time_s'] = round(float(time_centers[pi]),4)
        summary['peak_sig_distance'] = round(float(distances[pi]),6)
        summary['peak_sig_p_corrected'] = float(p_corr[pi])
        summary['peak_sig_p_uncorrected'] = float(p_values[pi])
    return summary

def compute_diagnostic_stats(pair, time_centers, scale=1/1000):
    diag = {}; dist = pair['distance'] * scale
    if np.all(np.isnan(dist)): return {'skipped': True}
    post_mask = time_centers > 0; pre_mask = time_centers < 0
    ci_width = (pair['ci_high'] - pair['ci_low']) * scale
    diag['ci_width_mean'] = round(float(np.nanmean(ci_width)),6)
    diag['ci_width_max'] = round(float(np.nanmax(ci_width)),6)
    diag['ci_width_mean_post'] = round(float(np.nanmean(ci_width[post_mask])),6) if np.any(post_mask) else None
    boot = pair['boot_distribution'] * scale; bm = np.nanmean(boot, axis=0); bs = np.nanstd(boot, axis=0)
    diag['boot_mean_post'] = round(float(np.nanmean(bm[post_mask])),6) if np.any(post_mask) else None
    diag['boot_std_mean_post'] = round(float(np.nanmean(bs[post_mask])),6) if np.any(post_mask) else None
    diag['boot_mean_over_observed'] = round(float(np.nanmean(bm[post_mask])/(np.nanmean(dist[post_mask])+1e-12)),4)
    if 'null_distribution' in pair:
        null = pair['null_distribution'] * scale; nm = np.nanmean(null, axis=0); ns = np.nanstd(null, axis=0)
        diag['null_mean_post'] = round(float(np.nanmean(nm[post_mask])),6) if np.any(post_mask) else None
        diag['null_std_mean_post'] = round(float(np.nanmean(ns[post_mask])),6) if np.any(post_mask) else None
        diag['null_max_post'] = round(float(np.nanmax(null[:,post_mask])),6) if np.any(post_mask) else None
        nsp = np.nanmean(ns[post_mask]) if np.any(post_mask) else 0
        if nsp > 0: diag['effect_to_null_ratio'] = round(float((np.nanmean(dist[post_mask])-np.nanmean(nm[post_mask]))/nsp),4)
        else: diag['effect_to_null_ratio'] = None
        if np.any(post_mask):
            diag['frac_null_exceeds_observed'] = round(float(np.mean(np.nanmean(null[:,post_mask],axis=1)>=np.nanmean(dist[post_mask]))),4)
    if np.any(pre_mask):
        diag['baseline_distance_mean'] = round(float(np.nanmean(dist[pre_mask])),6)
        diag['baseline_distance_std'] = round(float(np.nanstd(dist[pre_mask])),6)
        if np.any(post_mask): diag['post_minus_baseline'] = round(float(np.nanmean(dist[post_mask])-np.nanmean(dist[pre_mask])),6)
    return diag

# ============================================================================
# SIGNIFICANCE BAR HELPERS
# ============================================================================
def _draw_sig_bar(ax, bool_arr, timeline, y_val, color, linewidth, alpha):
    clusters = []; start = None
    for i, flag in enumerate(bool_arr):
        if flag and start is None: start = i
        elif not flag and start is not None: clusters.append((start, i-1)); start = None
    if start is not None: clusters.append((start, len(bool_arr)-1))
    for s, e in clusters:
        ax.plot([timeline[s], timeline[e]], [y_val, y_val], color=color, linewidth=linewidth,
                alpha=alpha, solid_capstyle='butt', label='_sig_line')

def add_sig_lines_two_tier(ax, fdr_arrays, cluster_arrays, colors, timeline,
                           linewidth_fdr=3, linewidth_cluster=5, alpha_fdr=0.35, alpha_cluster=0.8):
    y_range = np.ptp(ax.get_ylim()); spacing = y_range * 0.05; y_min = ax.get_ylim()[0]
    row = 0
    for t_idx in range(len(fdr_arrays)):
        color = colors[t_idx]
        row += 1; _draw_sig_bar(ax, fdr_arrays[t_idx], timeline, y_min+spacing*row, color, linewidth_fdr, alpha_fdr)
        row += 1; _draw_sig_bar(ax, cluster_arrays[t_idx], timeline, y_min+spacing*row, color, linewidth_cluster, alpha_cluster)

def add_sig_legend_annotation(ax, y_frac=0.02, fontsize=5):
    ax.text(0.0, y_frac, 'thin = FDR per-bin · thick = cluster-corrected',
            transform=ax.transAxes, ha='left', va='bottom', fontsize=fontsize, color='0.4', fontstyle='italic')

def _format_p(p):
    """Compact p-value formatting for in-panel annotations."""
    if p < 0.001: return f'{p:.1e}'
    return f'{p:.3f}'

# ============================================================================
# CONFIGURATION
# ============================================================================

with open(os.path.join(data_dir, 'metadata.yaml'), 'r') as f: metadata = yaml.safe_load(f)
area = metadata['area']; ci = metadata['ci']; z_scored = metadata['z_score_data']
y_lim_scale = 1/7 if z_scored else 1
include_opto = metadata.get('include_opto', True); p_alpha = 0.01; save_dir = data_dir  # stats files (yaml/xlsx) stay with the data
_area = metadata['area']; _area = ''.join(_area) if isinstance(_area, (list, tuple)) else str(_area)
fig_dir = os.path.join(config.FIGURE_DIR, 'dPCA', _area, 'euclidean_distance',
                       os.path.basename(os.path.normpath(data_dir)))
os.makedirs(fig_dir, exist_ok=True)

# ============================================================================
# FIGURE DEFINITIONS
# ============================================================================
FIGURES = [
    # ==== DECISION SUBSPACE (d) ====
    {'marg': 'd', 'fig_name': 'euclid_d_visual_LvsR', 'fig_title': 'Decision: L vs R separation — Visual trials',
     'y_lim': (-0.35, 3.5), 'x_lim': (-1, 1), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual L vs R (Control)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual L vs R (Ctrl)'},
            {'pair_label': 'Visual L vs R (Opto)', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Visual L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual L vs R (Control)', 'opto_label': 'Visual L vs R (Opto)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'}]}]},

    {'marg': 'd', 'fig_name': 'euclid_d_all_licks', 'fig_title': 'Ctrl vs opto in L, R trials',
     'y_lim': (-0.1, 1), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'All left licks: Ctrl vs Opto', 'color': 'blue', 'linestyle': '-',
             'label': 'Left licks opto deflection'},
            {'pair_label': 'All right licks: Ctrl vs Opto', 'color': 'cornflowerblue', 'linestyle': '--',
             'label': 'Right licks opto deflection', 'requires_opto': True}]}]},

    {'marg': 'd', 'fig_name': 'euclid_d_auditory_LvsR', 'fig_title': 'Decision: L vs R separation — Auditory trials',
     'y_lim': (-0.25, 2.5), 'x_lim': (-1, 1), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Auditory L vs R (Control)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory L vs R (Ctrl)'},
            {'pair_label': 'Auditory L vs R (Opto)', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Auditory L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Auditory L vs R (Control)', 'opto_label': 'Auditory L vs R (Opto)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},
    {'marg': 'd', 'fig_name': 'euclid_d_auditory_hit_vs_visual_hit', 'requires_opto': True,
     'fig_title': 'Decision: Difference in opto effects between visual and auditory hits', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Left: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Visual Opto Effect in hits'},
            {'pair_label': 'Auditory Right: Ctrl vs Opto', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Auditory Opto Effect in hits'}]},
        {'type': 'differences', 'title': 'Opto effect\n(Visual - Auditory)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Visual Left: Ctrl vs Opto', 'opto_label': 'Auditory Right: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Opto difference'}]}]},
    {'marg': 'd', 'fig_name': 'euclid_d_LvsR_combined', 'fig_title': 'Decision: L vs R separation — Visual vs Auditory',
     'y_lim': (-0.4, 4), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual L vs R (Control)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual L vs R (Ctrl)'},
            {'pair_label': 'Visual L vs R (Opto)', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Visual L vs R (Opto)', 'requires_opto': True},
            {'pair_label': 'Auditory L vs R (Control)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory L vs R (Ctrl)'},
            {'pair_label': 'Auditory L vs R (Opto)', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Auditory L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual L vs R (Control)', 'opto_label': 'Visual L vs R (Opto)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'},
            {'ctrl_label': 'Auditory L vs R (Control)', 'opto_label': 'Auditory L vs R (Opto)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},

    # --- NEW: All trials L vs R (ctrl & opto) in d-subspace ---
    {'marg': 'd', 'fig_name': 'euclid_d_all_trials_LvsR', 'fig_title': 'Decision: L vs R separation — All trials (Vis + Aud)',
     'y_lim': (-0.6, 2.5), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'All trials L vs R (Control)', 'color': 'black', 'linestyle': '-', 'label': 'All trials L vs R (Ctrl)'},
            {'pair_label': 'All trials L vs R (Opto)', 'color': 'grey', 'linestyle': '--', 'label': 'All trials L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'All trials L vs R (Control)', 'opto_label': 'All trials L vs R (Opto)', 'color': 'black', 'linestyle': '-', 'label': 'All trials'}]}]},

    # --- NEW: Visual hit vs Auditory hit separation in d-subspace (ctrl & opto) ---
    {'marg': 'd', 'fig_name': 'euclid_d_vis_vs_aud_hits', 'fig_title': 'Decision: Visual vs Auditory hit separation',
     'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual hit vs Auditory hit (Ctrl)', 'color': 'purple', 'linestyle': '-', 'label': 'Vis vs Aud hits (Ctrl)'},
            {'pair_label': 'Visual hit vs Auditory hit (Opto)', 'color': 'plum', 'linestyle': '--', 'label': 'Vis vs Aud hits (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual hit vs Auditory hit (Ctrl)', 'opto_label': 'Visual hit vs Auditory hit (Opto)', 'color': 'purple', 'linestyle': '-', 'label': 'Vis vs Aud hits'}]}]},

    {'marg': 'd', 'fig_name': 'euclid_d_saliency_scaling_visual', 'fig_title': 'Decision: saliency scaling',
     'requires_opto': False, 'y_lim': (-0.1, 1), 'x_lim': (-1, 0.1), 'panel_w_mm': 20, 'panel_h_mm': 20, 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (visual misses)', 'color': 'blue', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (visual hits)', 'color': 'blue', 'linestyle': '-', 'label': 'Hits'}]}]},
    {'marg': 'd', 'fig_name': 'euclid_d_saliency_scaling_auditory', 'fig_title': 'Decision: saliency scaling',
     'requires_opto': False, 'y_lim': (-0.1, 1), 'x_lim': (-1, 0.1), 'panel_w_mm': 20, 'panel_h_mm': 20, 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (Auditory misses)', 'color': 'red', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (Auditory hits)', 'color': 'red', 'linestyle': '-', 'label': 'Hits'}]}]},

    {'marg': 'd', 'fig_name': 'euclid_d_L_vis_vs_aud', 'fig_title': 'Decision: Vis vs aud during left licks',
     'panels': [
         {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
             {'pair_label': 'Visual vs Auditory L', 'color': 'blue', 'linestyle': '-',
              'label': 'Vis L vs Auditory L'},
             {'pair_label': 'Visual R vs Auditory R', 'color': 'red', 'linestyle': '--',
              'label': 'Vis R vs Auditory R', 'requires_opto': False}]}]},

    # ==== STIMULUS SUBSPACE (s) ====
{'marg': 's', 'fig_name': 'euclid_s_vis_vs_aud_all_responded', 'fig_title': 'Stimulus: Vis vs Aud separation — All responded trials',
 'panels': [
    {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
        {'pair_label': 'Vis vs Aud all responded (Control)', 'color': 'purple', 'linestyle': '-', 'label': 'Vis vs Aud all responded (Ctrl)'},
        {'pair_label': 'Vis vs Aud all responded (Opto)', 'color': 'plum', 'linestyle': '--', 'label': 'Vis vs Aud all responded (Opto)', 'requires_opto': True}]},
    {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
        {'ctrl_label': 'Vis vs Aud all responded (Control)', 'opto_label': 'Vis vs Aud all responded (Opto)', 'color': 'purple', 'linestyle': '-', 'label': 'All responded'}]}]},

    {'marg': 's', 'fig_name': 'euclid_s_hits', 'fig_title': 'Stimulus: Vis vs Aud separation — Hits', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Vis vs Aud Hits (Control)', 'color': 'purple', 'linestyle': '-', 'label': 'Vis vs Aud Hits (Ctrl)'},
            {'pair_label': 'Vis vs Aud Hits (Opto)', 'color': 'plum', 'linestyle': '--', 'label': 'Vis vs Aud Hits (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Hits (Control)', 'opto_label': 'Vis vs Aud Hits (Opto)', 'color': 'purple', 'linestyle': '-', 'label': 'Hits'}]}]},

    # --- NEW: Vis vs Aud Misses (ctrl & opto, with opto-effect panel) ---
    {'marg': 's', 'fig_name': 'euclid_s_misses', 'fig_title': 'Stimulus: Vis vs Aud separation — Misses', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Vis vs Aud Misses (Control)', 'color': 'darkgreen', 'linestyle': '-', 'label': 'Vis vs Aud Misses (Ctrl)'},
            {'pair_label': 'Vis vs Aud Misses (Opto)', 'color': 'limegreen', 'linestyle': '--', 'label': 'Vis vs Aud Misses (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Misses (Control)', 'opto_label': 'Vis vs Aud Misses (Opto)', 'color': 'darkgreen', 'linestyle': '-', 'label': 'Misses'}]}]},

    # --- NEW: Vis vs Aud Errors (ctrl & opto, with opto-effect panel) ---
    {'marg': 's', 'fig_name': 'euclid_s_errors', 'fig_title': 'Stimulus: Vis vs Aud separation — Errors', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Vis vs Aud Errors (Control)', 'color': 'darkorange', 'linestyle': '-', 'label': 'Vis vs Aud Errors (Ctrl)'},
            {'pair_label': 'Vis vs Aud Errors (Opto)', 'color': 'peachpuff', 'linestyle': '--', 'label': 'Vis vs Aud Errors (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Errors (Control)', 'opto_label': 'Vis vs Aud Errors (Opto)', 'color': 'darkorange', 'linestyle': '-', 'label': 'Errors'}]}]},

    # --- NEW: Within-modality Hit vs Error in s-subspace ---
    {'marg': 's', 'fig_name': 'euclid_s_within_modality_hit_vs_error',
     'fig_title': 'Stimulus: Hit vs Error within modality (s-subspace)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual hit vs error', 'color': 'blue', 'linestyle': '-', 'label': 'Visual hit vs error'},
            {'pair_label': 'Auditory hit vs error', 'color': 'red', 'linestyle': '-', 'label': 'Auditory hit vs error'}]}]},

    {'marg': 's', 'fig_name': 'euclid_s_left_licks', 'fig_title': 'Stimulus: Vis vs Aud separation — Left licks', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Vis vs Aud Left licks (Control)', 'color': 'darkgreen', 'linestyle': '-', 'label': 'Vis vs Aud Left (Ctrl)'},
            {'pair_label': 'Vis vs Aud Left licks (Opto)', 'color': 'lightgreen', 'linestyle': '--', 'label': 'Vis vs Aud Left (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Left licks (Control)', 'opto_label': 'Vis vs Aud Left licks (Opto)', 'color': 'darkgreen', 'linestyle': '-', 'label': 'Left licks'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_right_licks', 'fig_title': 'Stimulus: Vis vs Aud separation — Right licks', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Vis vs Aud Right licks (Control)', 'color': 'darkorange', 'linestyle': '-', 'label': 'Vis vs Aud Right (Ctrl)'},
            {'pair_label': 'Vis vs Aud Right licks (Opto)', 'color': 'peachpuff', 'linestyle': '--', 'label': 'Vis vs Aud Right (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Right licks (Control)', 'opto_label': 'Vis vs Aud Right licks (Opto)', 'color': 'darkorange', 'linestyle': '-', 'label': 'Right licks'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_auditory_hit_vs_visual_hit', 'requires_opto': True,
     'fig_title': 'Stimulus: Difference in opto effects between visual and auditory hits', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Left: Ctrl vs Opto', 'color': 'blue', 'linestyle': '-', 'label': 'Visual Opto Effect in hits'},
            {'pair_label': 'Auditory Right: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Auditory Opto Effect in hits'}]},
        {'type': 'differences', 'title': 'Opto effect\n(Visual - Auditory)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Visual Left: Ctrl vs Opto', 'opto_label': 'Auditory Right: Ctrl vs Opto', 'color': 'black', 'linestyle': '-', 'label': 'Opto difference'}]}]},

    # --- NEW: Visual and Auditory opto shift (all responded, combined L+R) ---
    {'marg': 's', 'fig_name': 'euclid_s_all_responded_opto_shift', 'requires_opto': True,
     'fig_title': 'Stimulus: Opto shift — all responded trials (combined L+R)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance\n(Ctrl vs Opto)', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual all responded: Ctrl vs Opto', 'color': 'blue', 'linestyle': '-', 'label': 'Visual (all responded)'},
            {'pair_label': 'Auditory all responded: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Auditory (all responded)'}]},
        {'type': 'differences', 'title': 'Opto effect\n(Visual − Auditory)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Visual all responded: Ctrl vs Opto', 'opto_label': 'Auditory all responded: Ctrl vs Opto',
             'color': 'black', 'linestyle': '-', 'label': 'Vis − Aud opto shift'}]}]},

    # --- NEW: Visual L vs R in stimulus subspace ---
    {'marg': 's', 'fig_name': 'euclid_s_visual_LvsR', 'requires_opto': True,
     'fig_title': 'Stimulus: Visual L vs R separation in s-subspace', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual L vs R (Control) [s]', 'color': 'blue', 'linestyle': '-', 'label': 'Visual L vs R (Ctrl)'},
            {'pair_label': 'Visual L vs R (Opto) [s]', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Visual L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual L vs R (Control) [s]', 'opto_label': 'Visual L vs R (Opto) [s]', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'}]}]},

    # --- NEW: Auditory L vs R in stimulus subspace ---
    {'marg': 's', 'fig_name': 'euclid_s_auditory_LvsR', 'requires_opto': True,
     'fig_title': 'Stimulus: Auditory L vs R separation in s-subspace', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Auditory L vs R (Control) [s]', 'color': 'red', 'linestyle': '-', 'label': 'Auditory L vs R (Ctrl)'},
            {'pair_label': 'Auditory L vs R (Opto) [s]', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Auditory L vs R (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Auditory L vs R (Control) [s]', 'opto_label': 'Auditory L vs R (Opto) [s]', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},
# --- NEW: Combined opto effect (Ctrl − Opto) for Visual & Auditory L vs R in s-subspace ---
    {'marg': 's', 'fig_name': 'euclid_s_visual_auditory_LvsR_optoeffect', 'requires_opto': True,
     'fig_title': 'Stimulus: Opto effect on L vs R separation — Visual vs Auditory', 'panels': [
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual L vs R (Control) [s]', 'opto_label': 'Visual L vs R (Opto) [s]', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'},
            {'ctrl_label': 'Auditory L vs R (Control) [s]', 'opto_label': 'Auditory L vs R (Opto) [s]', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},

    {'marg': 's', 'fig_name': 'euclid_s_combined_diff', 'requires_opto': True,
     'fig_title': 'Stimulus: Opto effect on Vis vs Aud separation', 'panels': [
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Vis vs Aud Hits (Control)', 'opto_label': 'Vis vs Aud Hits (Opto)', 'color': 'purple', 'linestyle': '-', 'label': 'Hits'},
            {'ctrl_label': 'Vis vs Aud Left licks (Control)', 'opto_label': 'Vis vs Aud Left licks (Opto)', 'color': 'darkgreen', 'linestyle': '--', 'label': 'Left licks'},
            {'ctrl_label': 'Vis vs Aud Right licks (Control)', 'opto_label': 'Vis vs Aud Right licks (Opto)', 'color': 'darkorange', 'linestyle': ':', 'label': 'Right licks'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_misses_hits_errors', 'requires_opto': False, 'y_lim': (-0.6, 6),
     'fig_title': 'Stimulus: Distance for hits, misses and errors', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Vis vs Aud Hits (Control)', 'color': 'purple', 'linestyle': '-', 'label': 'Hits'},
            {'pair_label': 'Vis vs Aud Misses (Control)', 'color': 'darkgreen', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Vis vs Aud Errors (Control)', 'color': 'darkorange', 'linestyle': '--', 'label': 'Errors'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_misses_opto', 'requires_opto': True,
     'fig_title': 'Stimulus: Distance misses (ctrl vs opto)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Aud miss ctrl vs opto)', 'color': 'red', 'linestyle': '-', 'label': 'Aud'},
            {'pair_label': 'Vis miss ctrl vs opto)', 'color': 'blue', 'linestyle': '-', 'label': 'Vis'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_saliency_scaling', 'requires_opto': False, 'y_lim': (-0.1, 2), 'x_lim': (-0.1, 1), 'panel_w_mm': 20,
     'fig_title': 'Stimulus: saliency scaling', 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (visual misses)', 'color': 'blue', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (visual hits)', 'color': 'blue', 'linestyle': '-', 'label': 'Hits'},
            {'pair_label': 'Saliency scaling (Auditory misses)', 'color': 'red', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (Auditory hits)', 'color': 'red', 'linestyle': '-', 'label': 'Hits'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_saliency_scaling_visual', 'requires_opto': False, 'y_lim': (-0.1, 1), 'x_lim': (-1, 0.5),
     'panel_w_mm': 20, 'panel_h_mm': 20, 'fig_title': 'Stimulus: saliency scaling', 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (visual misses)', 'color': 'blue', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (visual hits)', 'color': 'blue', 'linestyle': '-', 'label': 'Hits'}]}]},
    {'marg': 's', 'fig_name': 'euclid_s_saliency_scaling_auditory', 'requires_opto': False, 'y_lim': (-0.1, 2), 'x_lim': (-0.1, 1),
     'panel_w_mm': 20, 'panel_h_mm': 20, 'fig_title': 'Stimulus: saliency scaling', 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (Auditory misses)', 'color': 'red', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (Auditory hits)', 'color': 'red', 'linestyle': '-', 'label': 'Hits'}]}]},

    # ==== STIMULUS x DECISION (sd) ====
    {'marg': 'sd', 'fig_name': 'euclid_sd_visual_hit_vs_error', 'fig_title': 'Stim×Dec: Hit vs Error — Visual trials', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Hit vs Error (Control)', 'color': 'blue', 'linestyle': '-', 'label': 'Vis Hit vs Error (Ctrl)'},
            {'pair_label': 'Visual Hit vs Error (Opto)', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Vis Hit vs Error (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual Hit vs Error (Control)', 'opto_label': 'Visual Hit vs Error (Opto)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'}]}]},
    {'marg': 'sd', 'fig_name': 'euclid_sd_auditory_hit_vs_error', 'fig_title': 'Stim×Dec: Hit vs Error — Auditory trials', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Auditory Hit vs Error (Control)', 'color': 'red', 'linestyle': '-', 'label': 'Aud Hit vs Error (Ctrl)'},
            {'pair_label': 'Auditory Hit vs Error (Opto)', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Aud Hit vs Error (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Auditory Hit vs Error (Control)', 'opto_label': 'Auditory Hit vs Error (Opto)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},
    {'marg': 'sd', 'fig_name': 'euclid_sd_hit_vs_error_combined', 'fig_title': 'Stim×Dec: Hit vs Error — Visual vs Auditory',
     'y_lim': (-0.6, 6), 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Hit vs Error (Control)', 'color': 'blue', 'linestyle': '-', 'label': 'Vis Hit vs Err (Ctrl)'},
            {'pair_label': 'Visual Hit vs Error (Opto)', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Vis Hit vs Err (Opto)', 'requires_opto': True},
            {'pair_label': 'Auditory Hit vs Error (Control)', 'color': 'red', 'linestyle': '-', 'label': 'Aud Hit vs Err (Ctrl)'},
            {'pair_label': 'Auditory Hit vs Error (Opto)', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Aud Hit vs Err (Opto)', 'requires_opto': True}]},
        {'type': 'differences', 'title': 'Opto effect\n(Ctrl − Opto)', 'show_ci': True, 'hline_zero': True, 'requires_opto': True, 'traces': [
            {'ctrl_label': 'Visual Hit vs Error (Control)', 'opto_label': 'Visual Hit vs Error (Opto)', 'color': 'blue', 'linestyle': '-', 'label': 'Visual'},
            {'ctrl_label': 'Auditory Hit vs Error (Control)', 'opto_label': 'Auditory Hit vs Error (Opto)', 'color': 'red', 'linestyle': '-', 'label': 'Auditory'}]}]},

    # --- NEW: Hit trajectory shift (Ctrl vs Opto) for Visual and Auditory in sd-subspace ---
    {'marg': 'sd', 'fig_name': 'euclid_sd_hit_ctrl_vs_opto', 'requires_opto': True,
     'fig_title': 'Stim×Dec: Hit trajectory shift (Ctrl vs Opto)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance\n(Ctrl vs Opto)', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Hit: Ctrl vs Opto', 'color': 'blue', 'linestyle': '-', 'label': 'Visual hits'},
            {'pair_label': 'Auditory Hit: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Auditory hits'}]},
        {'type': 'differences', 'title': 'Opto shift\n(Visual − Auditory)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Visual Hit: Ctrl vs Opto', 'opto_label': 'Auditory Hit: Ctrl vs Opto',
             'color': 'black', 'linestyle': '-', 'label': 'Vis − Aud hit shift'}]}]},

    {'marg': 'sd', 'fig_name': 'euclid_sd_saliency_scaling', 'requires_opto': False,
     'fig_title': 'SD: saliency scaling', 'panels': [
        {'type': 'distances', 'title': 'Stimulus separation', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'pair_label': 'Saliency scaling (visual misses)', 'color': 'blue', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (visual hits)', 'color': 'blue', 'linestyle': '-', 'label': 'Hits'},
            {'pair_label': 'Saliency scaling (Auditory misses)', 'color': 'red', 'linestyle': ':', 'label': 'Misses'},
            {'pair_label': 'Saliency scaling (Auditory hits)', 'color': 'red', 'linestyle': '-', 'label': 'Hits'}]}]},

# --- NEW: Error trajectory shift (Ctrl vs Opto) for Visual and Auditory in sd-subspace ---
    {'marg': 'sd', 'fig_name': 'euclid_sd_error_ctrl_vs_opto', 'requires_opto': True,
     'fig_title': 'Stim×Dec: Error trajectory shift (Ctrl vs Opto)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance\n(Ctrl vs Opto)', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Error: Ctrl vs Opto', 'color': 'cornflowerblue', 'linestyle': '-', 'label': 'Visual errors'},
            {'pair_label': 'Auditory Error: Ctrl vs Opto', 'color': 'lightsalmon', 'linestyle': '-', 'label': 'Auditory errors'}]},
        {'type': 'differences', 'title': 'Opto shift\n(Visual − Auditory)', 'show_ci': True, 'hline_zero': True, 'traces': [
            {'ctrl_label': 'Visual Error: Ctrl vs Opto', 'opto_label': 'Auditory Error: Ctrl vs Opto',
             'color': 'black', 'linestyle': '-', 'label': 'Vis − Aud error shift'}]}]},

    # --- NEW (optional): Hit + Error trajectory shifts together, all four traces ---
    {'marg': 'sd', 'fig_name': 'euclid_sd_hit_error_ctrl_vs_opto_combined', 'requires_opto': True,
     'fig_title': 'Stim×Dec: Hit & Error trajectory shift (Ctrl vs Opto)', 'panels': [
        {'type': 'distances', 'title': 'Euclidean distance\n(Ctrl vs Opto)', 'show_ci': True, 'traces': [
            {'pair_label': 'Visual Hit: Ctrl vs Opto', 'color': 'blue', 'linestyle': '-', 'label': 'Visual hits'},
            {'pair_label': 'Auditory Hit: Ctrl vs Opto', 'color': 'red', 'linestyle': '-', 'label': 'Auditory hits'},
            {'pair_label': 'Visual Error: Ctrl vs Opto', 'color': 'cornflowerblue', 'linestyle': '--', 'label': 'Visual errors'},
            {'pair_label': 'Auditory Error: Ctrl vs Opto', 'color': 'lightsalmon', 'linestyle': '--', 'label': 'Auditory errors'}]}]},
]

# ============================================================================
# PLOTTING
# ============================================================================
print("Generating figures...")
sig_report = {}

for fig_spec in FIGURES:
    marg_key = fig_spec['marg']; fig_name = fig_spec['fig_name']; fig_title = fig_spec.get('fig_title', fig_name)
    if fig_spec.get('requires_opto', False) and not include_opto: print(f"  Skipping '{fig_name}' (requires opto)"); continue
    euclid_file = os.path.join(data_dir, f'euclidean_{marg_key}.npz')
    if not os.path.exists(euclid_file): print(f"  WARNING: {euclid_file} not found. Skipping '{fig_name}'."); continue
    time_centers, pairs = load_euclidean(data_dir, marg_key)
    label_to_idx = build_label_to_index(pairs)
    panels = [p for p in fig_spec['panels'] if not (p.get('requires_opto', False) and not include_opto)]
    if not panels: print(f"  Skipping '{fig_name}' (no panels remain)"); continue
    dw = SUBPLOT_W / MM_TO_IN; dh = SUBPLOT_H / MM_TO_IN
    panel_sizes = [(p.get('panel_w_mm', fig_spec.get('panel_w_mm', dw))*MM_TO_IN,
                    p.get('panel_h_mm', fig_spec.get('panel_h_mm', dh))*MM_TO_IN) for p in panels]
    fig, axes_1d = make_figure_with_variable_panels(panel_sizes); n_rows = len(panels)

    for row_idx, panel in enumerate(panels):
        ax = axes_1d[row_idx]; pt = panel['type']; show_ci = panel.get('show_ci', True); ci_alpha = 0.15
        traces = [t for t in panel['traces'] if not (t.get('requires_opto', False) and not include_opto)]

        if pt == 'distances':
            fdr_arr, clust_arr, sig_col = [], [], []; has_clust = False
            # Collect per-trace window info for later shading + annotation
            trace_window_info = []   # list of dicts: {color, p_win, obs_win, window, label}
            for trace in traces:
                pl = trace.get('pair_label')
                if pl is None or pl not in label_to_idx:
                    if pl is not None: print(f"    WARNING: pair '{pl}' not found for {marg_key}.")
                    continue
                pair = pairs[label_to_idx[pl]]; lbl = trace.get('label', pair['label'])
                na, nb = pair['n_trials_A'], pair['n_trials_B']
                ax.plot(time_centers, pair['distance']/1000, color=trace['color'],
                        linestyle=trace.get('linestyle','-'), linewidth=1.5, label=f"{lbl} (n={na},{nb})")
                if show_ci: ax.fill_between(time_centers, pair['ci_low']/1000, pair['ci_high']/1000,
                                            color=trace['color'], alpha=ci_alpha, linewidth=0)
                if 'p_values' in pair:
                    rej, _, _, _ = multipletests(pair['p_values'], alpha=p_alpha, method='fdr_bh'); fdr_arr.append(rej)
                else: fdr_arr.append(np.zeros(len(time_centers), dtype=bool))
                if 'cluster_mask' in pair: clust_arr.append(pair['cluster_mask']); has_clust = True
                else: clust_arr.append(np.zeros(len(time_centers), dtype=bool))
                sig_col.append(trace['color'])
                # Record window info for this trace if present
                if 'p_value_window_avg' in pair:
                    trace_window_info.append({
                        'color': trace['color'], 'p_win': pair['p_value_window_avg'],
                        'obs_win': pair['observed_window_avg'] / 1000,  # rescale to plot units
                        'window': tuple(pair['window']), 'label': lbl, 'pair_label': pl,
                    })

            # --- Window shading: draw once per unique window (all traces usually share it) ---
            if trace_window_info:
                unique_windows = list({tw['window'] for tw in trace_window_info})
                for win in unique_windows:
                    ax.axvspan(win[0], win[1], alpha=0.06, color='darkred', zorder=0)

            if fdr_arr:
                if has_clust: add_sig_lines_two_tier(ax, fdr_arr, clust_arr, sig_col, timeline=time_centers); add_sig_legend_annotation(ax)
                else: plot_utils.add_sig_lines_ax_editable(ax, fdr_arr, colors=sig_col, timeline=time_centers)
                ns_off = 0
                for fa, ca, sc in zip(fdr_arr, clust_arr, sig_col):
                    if not np.any(fa) and not np.any(ca):
                        ax.text(0.97, 0.95-ns_off, 'ns', transform=ax.transAxes, ha='right', va='top',
                                fontsize=FONT_DEFAULT, color=sc, fontstyle='italic'); ns_off += 0.1
                for trace in traces:
                    pl = trace.get('pair_label')
                    if pl is None or pl not in label_to_idx: continue
                    pair = pairs[label_to_idx[pl]]
                    if 'p_values' not in pair: continue
                    summary = summarize_significance(time_centers, pair['p_values'], pair['distance']/1000, alpha=p_alpha)
                    summary['n_trials_A'] = pair['n_trials_A']; summary['n_trials_B'] = pair['n_trials_B']
                    summary['diagnostics'] = compute_diagnostic_stats(pair, time_centers, scale=1/1000)
                    # --- Window-averaged test results into YAML (if available) ---
                    if 'p_value_window_avg' in pair:
                        summary['window_test'] = {
                            'window': [float(pair['window'][0]), float(pair['window'][1])],
                            'observed_window_avg': round(float(pair['observed_window_avg']) / 1000, 6),
                            'p_value_window_avg': float(pair['p_value_window_avg']),
                            'significant': bool(pair['p_value_window_avg'] < p_alpha),
                            'alpha': p_alpha,
                        }
                    if 'cluster_p_value' in pair and len(pair['cluster_p_value']) > 0:
                        summary['cluster_test'] = {'n_clusters': int(len(pair['cluster_p_value'])),
                            'clusters': [{'t_start': round(float(pair['cluster_t_start'][k]),4),
                                          't_end': round(float(pair['cluster_t_end'][k]),4),
                                          'mass': round(float(pair['cluster_mass'][k]),4),
                                          'p_value': float(pair['cluster_p_value'][k])}
                                         for k in range(len(pair['cluster_p_value']))],
                            'n_sig_clusters': int(np.sum(pair['cluster_p_value'] < p_alpha))}
                    else: summary['cluster_test'] = {'n_clusters': 0}
                    if fig_name not in sig_report:
                        sig_report[fig_name] = {'title': fig_title, 'marginalization': marg_key, 'area': area, 'panels': {}}
                    ptitle = panel.get('title', f'panel_{row_idx}')
                    if ptitle not in sig_report[fig_name]['panels']: sig_report[fig_name]['panels'][ptitle] = {}
                    sig_report[fig_name]['panels'][ptitle][pl] = summary

            # --- Window-averaged p-values: stacked color-coded annotations, bottom-right ---
            if trace_window_info:
                y_base = 0.05
                y_step = 0.08
                # Top label indicating what the window is
                win_for_header = trace_window_info[0]['window']
                header = f'window [{win_for_header[0]:.2f}, {win_for_header[1]:.2f}]s:'
                ax.text(0.97, y_base + y_step * len(trace_window_info), header,
                        transform=ax.transAxes, ha='right', va='bottom',
                        fontsize=5, color='darkred', fontstyle='italic')
                for k, tw in enumerate(trace_window_info):
                    star = '*' if tw['p_win'] < p_alpha else ''
                    txt = f"p_{{win}} = {_format_p(tw['p_win'])}{star}"
                    ax.text(0.97, y_base + y_step * k, txt,
                            transform=ax.transAxes, ha='right', va='bottom',
                            fontsize=FONT_DEFAULT, color=tw['color'], fontstyle='italic')

        elif pt == 'differences':
            for trace in traces:
                if 'diff_index' in trace:
                    _, diffs = load_euclidean_diff(data_dir, marg_key); d = diffs[trace['diff_index']]
                    dd = {'diff_distance': d['diff_distance'], 'diff_ci_low': d['diff_ci_low'], 'diff_ci_high': d['diff_ci_high']}
                else:
                    cl, ol = trace['ctrl_label'], trace['opto_label']
                    if cl not in label_to_idx or ol not in label_to_idx:
                        print(f"    WARNING: '{cl}' or '{ol}' not found for {marg_key}."); continue
                    dd = compute_diff_from_pairs(pairs[label_to_idx[cl]], pairs[label_to_idx[ol]], ci=ci)
                lbl = trace.get('label', 'Difference')
                ax.plot(time_centers, dd['diff_distance']/1000, color=trace['color'],
                        linestyle=trace.get('linestyle','-'), linewidth=1.5, label=lbl)
                if show_ci: ax.fill_between(time_centers, dd['diff_ci_low']/1000, dd['diff_ci_high']/1000,
                                            color=trace['color'], alpha=ci_alpha, linewidth=0)

        ax.axvline(0, color='k', linestyle='--', alpha=0.5, linewidth=0.8)
        if panel.get('hline_zero', False): ax.axhline(0, color='k', linestyle=':', alpha=0.5, linewidth=0.8)
        ax.grid(False)
        ylim = panel.get('y_lim', fig_spec.get('y_lim', None))
        if ylim is not None: ax.set_ylim(tuple(np.array(ylim)*y_lim_scale))
        xlim = panel.get('x_lim', fig_spec.get('x_lim', X_LIM_DEFAULT))
        if xlim is not None: ax.set_xlim(xlim)
        ax.tick_params(axis='both', which='major', labelsize=FONT_DEFAULT, length=3, width=0.6)
        ax.set_ylabel(panel.get('title',''), fontsize=FONT_DEFAULT)
        if row_idx == n_rows - 1: ax.set_xlabel('Time from stimulus (s)', fontsize=FONT_DEFAULT)
        ax.legend(fontsize=FONT_LEGEND, loc='best', frameon=False)
        if row_idx == 0: ax.set_title(f"{marg_key} subspace", fontsize=FONT_TITLE, fontweight='bold')

    fig.text(0.5, 1.0-(TOP_MARGIN*0.35)/fig.get_size_inches()[1], f"{fig_title} — {', '.join(area)}",
             ha='center', va='top', fontsize=FONT_SUPTITLE, fontweight='bold')
    fig.savefig(os.path.join(fig_dir, f'{fig_name}.png'), dpi=300, bbox_inches='tight')
    fig.savefig(os.path.join(fig_dir, f'{fig_name}.svg'), dpi=300, bbox_inches='tight')
    print(f"  Saved: {fig_name}.png / .svg"); plt.close(fig)

# ============================================================================
# SAVE SIGNIFICANCE SUMMARIES
# ============================================================================
sig_yaml_path = os.path.join(save_dir, 'significance_summaries.yaml')
class IndentDumper(yaml.Dumper):
    def increase_indent(self, flow=False, indentless=False): return super().increase_indent(flow, False)
def sci_float_representer(dumper, value):
    if value == 0.0: text = f'{value:.2e}'
    elif abs(value) < 0.001: text = f'{value:.4e}'
    elif abs(value) < 1: text = f'{value:.6f}'
    else: text = f'{value:.4f}'
    return dumper.represent_scalar('tag:yaml.org,2002:float', text)
IndentDumper.add_representer(float, sci_float_representer)
with open(sig_yaml_path, 'w') as f:
    yaml.dump(sig_report, f, Dumper=IndentDumper, default_flow_style=False, sort_keys=False, width=120)

# ============================================================================
# OPTO PERMUTATION TEST FIGURES
# ============================================================================
def load_opto_permutation(data_dir, marg_key):
    path = os.path.join(data_dir, f'opto_permutation_{marg_key}.npz')
    if not os.path.exists(path): return None, []
    data = np.load(path, allow_pickle=True); tc = data['time_centers']; comps = []; i = 0
    while f'comp{i}_label' in data:
        c = {'label': str(data[f'comp{i}_label']), 'type': str(data[f'comp{i}_type']),
             'observed': data[f'comp{i}_observed'], 'null_distribution': data[f'comp{i}_null_distribution'],
             'p_values': data[f'comp{i}_p_values']}
        if f'comp{i}_p_value_window_avg' in data:
            c['p_value_window_avg'] = float(data[f'comp{i}_p_value_window_avg'])
            c['observed_window_avg'] = float(data[f'comp{i}_observed_window_avg'])
            c['window'] = data[f'comp{i}_window']
        if f'comp{i}_cluster_mask' in data:
            c['cluster_mask'] = data[f'comp{i}_cluster_mask'].astype(bool)
            c['cluster_p_value'] = data[f'comp{i}_cluster_p_value']
            c['cluster_t_start'] = data[f'comp{i}_cluster_t_start']
            c['cluster_t_end'] = data[f'comp{i}_cluster_t_end']
            c['cluster_mass'] = data[f'comp{i}_cluster_mass']
        comps.append(c); i += 1
    return tc, comps

if include_opto:
    print("\nGenerating opto permutation figures...")
    for marg_key in ['d', 's', 'sd']:
        tc_op, comparisons = load_opto_permutation(data_dir, marg_key)
        if tc_op is None or len(comparisons) == 0: print(f"  No opto permutation data for '{marg_key}'."); continue
        nc = len(comparisons); fn = f'opto_permutation_{marg_key}'
        ps = [(SUBPLOT_W, 15*MM_TO_IN)] * nc
        fig, axes_1d = make_figure_with_variable_panels(ps)
        for idx, comp in enumerate(comparisons):
            ax = axes_1d[idx]; obs = comp['observed']; pv = comp['p_values']; tt = comp['type']
            rej, pc, _, _ = multipletests(pv, alpha=p_alpha, method='fdr_bh')
            col = 'black'
            ax.plot(tc_op, obs/1000, color=col, linewidth=1.2)
            ax.axhline(0, color='k', linestyle=':', alpha=0.4, linewidth=0.6)
            ax.axvline(0, color='k', linestyle='--', alpha=0.4, linewidth=0.6)
            hc = 'cluster_mask' in comp
            fa = [rej]; ca = [comp['cluster_mask'] if hc else np.zeros_like(rej)]
            if hc: add_sig_lines_two_tier(ax, fa, ca, [col], timeline=tc_op); add_sig_legend_annotation(ax)
            else: plot_utils.add_sig_lines_ax(ax, [rej], colors=[col], timeline=tc_op)
            if not np.any(rej) and not (hc and np.any(comp['cluster_mask'])):
                ax.text(0.97, 0.95, 'ns', transform=ax.transAxes, ha='right', va='top', fontsize=FONT_DEFAULT, color=col, fontstyle='italic')
            if 'p_value_window_avg' in comp:
                pw = comp['p_value_window_avg']; win = comp['window']
                ps_str = f'p_{{win}} = {pw:.1e}' if pw < 0.001 else f'p_{{win}} = {pw:.3f}'
                ax.text(0.97, 0.05, ps_str, transform=ax.transAxes, ha='right', va='bottom', fontsize=FONT_DEFAULT, color='darkred', fontstyle='italic')
                ax.axvspan(win[0], win[1], alpha=0.06, color='darkred', zorder=0)
            if hc and len(comp.get('cluster_p_value', [])) > 0:
                for k in range(len(comp['cluster_p_value'])):
                    cp = comp['cluster_p_value'][k]
                    if cp < p_alpha:
                        tm = (comp['cluster_t_start'][k]+comp['cluster_t_end'][k])/2; yt = ax.get_ylim()[1]
                        cs = f'cluster p={cp:.1e}' if cp < 0.001 else f'cluster p={cp:.3f}'
                        ax.text(tm, yt*0.92, cs, ha='center', va='top', fontsize=5, color='darkblue', fontstyle='italic')
            ts = tt.replace('_',' ')
            ax.set_title(f"{comp['label']}  ({ts})", fontsize=FONT_DEFAULT, fontweight='bold', pad=3)
            ax.set_ylabel('Δ distance', fontsize=FONT_DEFAULT)
            ax.tick_params(axis='both', labelsize=FONT_DEFAULT, length=2, width=0.5);
            ax.grid(False)
            ax.set_xlim(X_LIM_DEFAULT)
            if idx == nc - 1: ax.set_xlabel('Time from stimulus (s)', fontsize=FONT_DEFAULT)
            summary = summarize_significance(tc_op, pv, np.abs(obs)/1000, alpha=p_alpha)
            if 'p_value_window_avg' in comp:
                summary['p_value_window_avg'] = comp['p_value_window_avg']
                summary['observed_window_avg'] = comp['observed_window_avg']
                summary['window'] = [float(x) for x in comp['window']]
            od = {}; null = comp['null_distribution']/1000; ob = obs/1000; pm = tc_op > 0
            nm = np.nanmean(null, axis=0); nsd = np.nanstd(null, axis=0)
            if np.any(pm):
                od['observed_mean_post'] = round(float(np.nanmean(ob[pm])),6)
                od['null_mean_post'] = round(float(np.nanmean(nm[pm])),6)
                od['null_std_mean_post'] = round(float(np.nanmean(nsd[pm])),6)
                od['null_max_post'] = round(float(np.nanmax(null[:,pm])),6)
                nsp = np.nanmean(nsd[pm])
                if nsp > 0: od['effect_to_null_ratio'] = round(float((np.nanmean(ob[pm])-np.nanmean(nm[pm]))/nsp),4)
                od['frac_null_exceeds_observed'] = round(float(np.mean(np.nanmean(null[:,pm],axis=1)>=np.nanmean(ob[pm]))),4)
            prm = tc_op < 0
            if np.any(prm): od['observed_mean_pre'] = round(float(np.nanmean(ob[prm])),6); od['null_mean_pre'] = round(float(np.nanmean(nm[prm])),6)
            summary['diagnostics'] = od
            if hc and len(comp.get('cluster_p_value',[])) > 0:
                summary['cluster_test'] = {'n_clusters': int(len(comp['cluster_p_value'])),
                    'clusters': [{'t_start': round(float(comp['cluster_t_start'][k]),4), 't_end': round(float(comp['cluster_t_end'][k]),4),
                                  'mass': round(float(comp['cluster_mass'][k]),4), 'p_value': float(comp['cluster_p_value'][k])}
                                 for k in range(len(comp['cluster_p_value']))],
                    'n_sig_clusters': int(np.sum(comp['cluster_p_value'] < p_alpha))}
            else: summary['cluster_test'] = {'n_clusters': 0}
            if fn not in sig_report: sig_report[fn] = {'title': f'Opto permutation tests — {marg_key}', 'marginalization': marg_key, 'area': area, 'panels': {}}
            sig_report[fn]['panels'][comp['label']] = summary
        fig.text(0.5, 1.0-(TOP_MARGIN*0.35)/fig.get_size_inches()[1],
                 f"Opto permutation tests — {marg_key} subspace — {', '.join(area)}",
                 ha='center', va='top', fontsize=FONT_SUPTITLE, fontweight='bold')
        fig.savefig(os.path.join(fig_dir, f'{fn}.png'), dpi=300, bbox_inches='tight')
        fig.savefig(os.path.join(fig_dir, f'{fn}.svg'), dpi=300, bbox_inches='tight')
        print(f"  Saved: {fn}.png / .svg ({nc} comparisons)"); plt.close(fig)
    with open(sig_yaml_path, 'w') as f:
        yaml.dump(sig_report, f, Dumper=IndentDumper, default_flow_style=False, sort_keys=False, width=120)
    print(f"  Updated significance summaries: {sig_yaml_path}")

print(f"  Saved significance summaries: {sig_yaml_path}")
print(f"\nDone! Plots saved to: {fig_dir}")

# ============================================================================
# DROP-IN SNIPPET: EXPORT RESULTS TO SPREADSHEET (3 tabs, cluster-only)
#   - Decision_subspace
#   - Stimulus_subspace
#   - Opto_permutation_tests
#
# Place at the END of your script (after sig_report has been fully populated).
# Requirements: pip install openpyxl
# ============================================================================
import os
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ----------------------------------------------------------------------------
# Contents of each tab
# ----------------------------------------------------------------------------
DECISION_BLOCKS = [
    ('Saliency scaling — visual',                          'euclid_d_saliency_scaling_visual',    None),
    ('Saliency scaling — auditory',                        'euclid_d_saliency_scaling_auditory',  None),
    ('Opto-effect difference: visual vs auditory hits',    'euclid_d_auditory_hit_vs_visual_hit', None),
]

STIMULUS_BLOCKS = [
    ('Saliency scaling (visual + auditory)',               'euclid_s_saliency_scaling',                 None),
    ('Vis vs Aud distance — hits, misses, errors',         'euclid_s_misses_hits_errors',               None),
    ('Hit vs Error within modality',                       'euclid_s_within_modality_hit_vs_error',     None),
    ('Opto shift — all responded trials (combined L+R)',   'euclid_s_all_responded_opto_shift',         None),
    ('Opto-effect difference: visual vs auditory hits',    'euclid_s_auditory_hit_vs_visual_hit',       None),
    ('Auditory L vs R separation',                         'euclid_s_auditory_LvsR',                    None),
    ('Visual L vs R separation',                           'euclid_s_visual_LvsR',                      None),
    ('Vis vs Aud separation — Misses (ctrl & opto)',       'euclid_s_misses',                           None),
    ('Miss trial trajectory shift (ctrl vs opto)',         'euclid_s_misses_opto',                      None),
]

PERM_FIG_NAMES = ['opto_permutation_d', 'opto_permutation_s', 'opto_permutation_sd']
PERM_SECTION_LABELS = {
    'opto_permutation_d':  'Decision subspace (d)',
    'opto_permutation_s':  'Stimulus subspace (s)',
    'opto_permutation_sd': 'Stim×Dec subspace (sd)',
}

# ----------------------------------------------------------------------------
# Styling
# ----------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
BLOCK_FILL  = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
BLOCK_FONT  = Font(bold=True, color='FFFFFF', size=11)
PANEL_FILL  = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
PANEL_FONT  = Font(bold=True, size=10)
SIG_FILL    = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
NS_FILL     = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
THIN        = Side(border_style='thin', color='BFBFBF')
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER      = Alignment(horizontal='center', vertical='center')
LEFT        = Alignment(horizontal='left',  vertical='center')

def _fmt_p(p):
    if p is None or (isinstance(p, float) and np.isnan(p)): return ''
    p = float(p)
    if p < 0.001: return f'{p:.2e}'
    return f'{p:.4f}'

def _set_header_row(ws, row, headers):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER

def _set_block_row(ws, row, label, span):
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = BLOCK_FONT; cell.fill = BLOCK_FILL; cell.alignment = LEFT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = 20

def _set_panel_row(ws, row, label, span):
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = PANEL_FONT; cell.fill = PANEL_FILL; cell.alignment = LEFT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)

def _autosize(ws, max_width=44):
    for col_cells in ws.columns:
        col_letter = None; max_len = 0
        for cell in col_cells:
            if not hasattr(cell, 'column_letter') or cell.column_letter is None: continue
            col_letter = cell.column_letter
            v = cell.value
            if v is None: continue
            l = len(str(v))
            if l > max_len: max_len = l
        if col_letter:
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), max_width)

# ----------------------------------------------------------------------------
# Row builders (cluster-only)
# ----------------------------------------------------------------------------
PAIR_HEADERS = [
    'Figure', 'Panel', 'Pair label', 'n_A', 'n_B',
    'Mean dist (pre)', 'Mean dist (post)', 'Peak dist', 'Peak time (s)',
    'n clusters', 'n sig clusters', 'Cluster p-values', 'Cluster windows (s)',
    'Cluster masses', 'sig (cluster)',
]

def _cluster_summary(clust, alpha):
    """Return (n, n_sig, p_str, win_str, mass_str, sig_flag) from a cluster_test dict."""
    if not clust or clust.get('n_clusters', 0) == 0:
        return 0, 0, '', '', '', 'no'
    cs = clust.get('clusters', []) or []
    p_str    = '; '.join(_fmt_p(c.get('p_value')) for c in cs)
    win_str  = '; '.join(f"[{c.get('t_start','')},{c.get('t_end','')}]" for c in cs)
    mass_str = '; '.join(f"{c.get('mass',''):g}" if isinstance(c.get('mass'), (int, float)) else str(c.get('mass','')) for c in cs)
    n_sig = clust.get('n_sig_clusters', 0)
    return clust.get('n_clusters', 0), n_sig, p_str, win_str, mass_str, ('yes' if n_sig and n_sig > 0 else 'no')

def _row_for_pair(fig_name, panel_title, pair_label, summary, alpha):
    if summary is None:
        return [fig_name, panel_title, pair_label] + [''] * (len(PAIR_HEADERS) - 3)
    clust = summary.get('cluster_test', {}) or {}
    n_cl, n_sig, p_str, win_str, mass_str, sig_flag = _cluster_summary(clust, alpha)
    return [
        fig_name, panel_title, pair_label,
        summary.get('n_trials_A', ''), summary.get('n_trials_B', ''),
        summary.get('mean_distance_pre', ''), summary.get('mean_distance_post', ''),
        summary.get('peak_distance', ''), summary.get('peak_time_s', ''),
        n_cl, n_sig, p_str, win_str, mass_str, sig_flag,
    ]

PERM_HEADERS = [
    'Subspace', 'Comparison',
    'n clusters', 'n sig clusters', 'Cluster p-values', 'Cluster windows (s)',
    'Cluster masses', 'sig (cluster)',
    'Observed mean (post)', 'Null mean (post)', 'Null std (post)',
    'Effect/Null ratio', 'Frac null >= obs',
]

def _row_for_perm(subspace_label, comp_label, summary, alpha):
    diag  = summary.get('diagnostics', {}) or {}
    clust = summary.get('cluster_test', {}) or {}
    n_cl, n_sig, p_str, win_str, mass_str, sig_flag = _cluster_summary(clust, alpha)
    return [
        subspace_label, comp_label,
        n_cl, n_sig, p_str, win_str, mass_str, sig_flag,
        diag.get('observed_mean_post', ''),
        diag.get('null_mean_post', ''),
        diag.get('null_std_mean_post', ''),
        diag.get('effect_to_null_ratio', ''),
        diag.get('frac_null_exceeds_observed', ''),
    ]

# ----------------------------------------------------------------------------
# Tab writers
# ----------------------------------------------------------------------------
def _color_sig_cell(ws, row, headers):
    if 'sig (cluster)' not in headers: return
    col_idx = headers.index('sig (cluster)') + 1
    v = ws.cell(row=row, column=col_idx).value
    if   v == 'yes': ws.cell(row=row, column=col_idx).fill = SIG_FILL
    elif v == 'no':  ws.cell(row=row, column=col_idx).fill = NS_FILL

def write_pair_tab(ws, tab_title, blocks, sig_report, alpha, area):
    span = len(PAIR_HEADERS)
    r = 1
    ws.cell(row=r, column=1, value=tab_title).font = Font(bold=True, size=14); r += 1
    area_str = ', '.join(area) if isinstance(area, (list, tuple)) else str(area)
    ws.cell(row=r, column=1, value=f"area: {area_str}   |   alpha = {alpha}   |   cluster-based permutation test only   |   distances rescaled by 1/1000").font = Font(italic=True, size=9); r += 2

    for section_label, fig_name, panel_filter in blocks:
        _set_block_row(ws, r, section_label, span); r += 1

        if fig_name not in sig_report:
            cell = ws.cell(row=r, column=1, value=f"[no data] '{fig_name}' not present in sig_report")
            cell.font = Font(italic=True, color='990000')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            r += 2; continue

        info   = sig_report[fig_name]
        panels = info.get('panels', {}) or {}
        for panel_title, pair_dict in panels.items():
            if panel_filter is not None and panel_title != panel_filter: continue
            _set_panel_row(ws, r, f"Panel: {panel_title}", span); r += 1
            _set_header_row(ws, r, PAIR_HEADERS); r += 1
            for pair_label, summary in pair_dict.items():
                row_vals = _row_for_pair(fig_name, panel_title, pair_label, summary, alpha)
                for c, v in enumerate(row_vals, start=1):
                    cell = ws.cell(row=r, column=c, value=v); cell.border = BORDER
                    cell.alignment = LEFT if c <= 3 else CENTER
                _color_sig_cell(ws, r, PAIR_HEADERS)
                r += 1
            r += 1
        r += 1
    _autosize(ws)

def write_perm_tab(ws, tab_title, fig_names, sig_report, alpha, area):
    span = len(PERM_HEADERS)
    r = 1
    ws.cell(row=r, column=1, value=tab_title).font = Font(bold=True, size=14); r += 1
    area_str = ', '.join(area) if isinstance(area, (list, tuple)) else str(area)
    ws.cell(row=r, column=1, value=f"area: {area_str}   |   alpha = {alpha}   |   cluster-based permutation test only").font = Font(italic=True, size=9); r += 2

    _set_header_row(ws, r, PERM_HEADERS); r += 1

    for fig_name in fig_names:
        if fig_name not in sig_report:
            cell = ws.cell(row=r, column=1, value=f"[no data] '{fig_name}' not present in sig_report")
            cell.font = Font(italic=True, color='990000')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            r += 2; continue

        subspace_label = PERM_SECTION_LABELS.get(fig_name, fig_name)
        _set_panel_row(ws, r, f"{subspace_label}  ({fig_name})", span); r += 1

        info   = sig_report[fig_name]
        panels = info.get('panels', {}) or {}
        for comp_label, summary in panels.items():
            row_vals = _row_for_perm(subspace_label, comp_label, summary, alpha)
            for c, v in enumerate(row_vals, start=1):
                cell = ws.cell(row=r, column=c, value=v); cell.border = BORDER
                cell.alignment = LEFT if c <= 2 else CENTER
            _color_sig_cell(ws, r, PERM_HEADERS)
            r += 1
        r += 1
    _autosize(ws)

# ----------------------------------------------------------------------------
# Build the workbook
# ----------------------------------------------------------------------------
xlsx_path = os.path.join(save_dir, 'results_summary.xlsx')
wb = Workbook()
wb.remove(wb.active)

ws_dec  = wb.create_sheet('Decision_subspace')
write_pair_tab(ws_dec,  'Decision subspace (d)', DECISION_BLOCKS,  sig_report, p_alpha, area)

ws_stim = wb.create_sheet('Stimulus_subspace')
write_pair_tab(ws_stim, 'Stimulus subspace (s)', STIMULUS_BLOCKS, sig_report, p_alpha, area)

ws_perm = wb.create_sheet('Opto_permutation_tests')
write_perm_tab(ws_perm, 'Opto permutation tests (all subspaces)', PERM_FIG_NAMES, sig_report, p_alpha, area)

wb.save(xlsx_path)
print(f"\nSaved spreadsheet: {xlsx_path}")